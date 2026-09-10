# -*- coding: utf-8 -*-
"""Importa a receita das 9 misturas feitas na casa (aba "Sub-Receitas" da
planilha de CMV do Artesanos): Tempero Smash, Tempero Batata, Maionese da
Casa, Molho Especial, Massa de Dadinho, Cebola Crispy, Bacon Empanado,
Sweet Barbecue e Cebola Caramelizada.

Pra que serve: com a receita cadastrada, o custo da mistura passa a ser
calculado (e se corrige sozinho quando o preço de um ingrediente muda), e
vender um produto que leva a mistura desconta os ingredientes do estoque —
sal, páprica, açúcar —, que é o que a casa compra. Pedido da Julia em
09/09/2026, pra planilha sair de cena.

Receita dentro de receita: o Molho Especial leva 2,8 kg de Maionese da
Casa, que tem receita própria. O sistema resolve em cascata.

O que não entra completo, e por quê:
- "Limão (suco)" da Maionese: a receita pede gramas de suco, o estoque
  conta a fruta (por unidade, ou por quilo na VMarket). Converter exigiria
  saber quanto suco dá um limão, então a linha entra com a quantidade em
  branco (e a receita fica incompleta até alguém preencher) em vez de
  sumir.
- Na planilha, toda mistura rende exatamente a soma dos ingredientes. Pra
  tempero é isso mesmo; cebola caramelizada, cebola crispy e bacon
  empanado vão ao fogo e rendem menos. O rendimento entra como está na
  planilha e fica marcado pra conferir na tela (Estoque → receita).

Ingrediente que já existe é reaproveitado mesmo com outro nome — o da
VMarket, que é o catálogo do Artesanos em produção ("Açúcar Refinado
1Kg"), ou com o tamanho do pacote no nome — pelas mesmas regras dos passos
do Açaí e da Tradiça (localizar_insumo). Se ele é contado por embalagem
(galão, pacote), a quantidade vira fração da embalagem pelo conteúdo
cadastrado; sem conteúdo, a linha entra em branco. Só o que não for achado
de jeito nenhum é criado, vinculado ao Artesanos, com o custo unitário da
própria sub-receita — e o relatório avisa, com os parecidos do cadastro,
porque criar um insumo que já existe partiria o estoque dele em dois.

Idempotente: mistura que já tem receita não é tocada (pode ter sido
ajustada na tela).

Uso:
  python importar_sub_receitas.py [planilha.xlsx]            # simulação
  python importar_sub_receitas.py [planilha.xlsx] --apply    # grava
"""
import sys
from datetime import datetime

import openpyxl

from backend.armazenamento import (
    conexao,
    criar_insumo,
    definir_receita_insumo,
    inicializar_banco,
    mapa_receita_insumo,
    _normalizar_nome_insumo,
)
from backend.nomes_insumo import localizar_insumo, nome_base
from importar_custos_insumo import _unidade
from importar_ficha_tecnica_faltante import quantidade_para_insumo

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass

PLANILHA_PADRAO = r"C:\Users\Guilherme\Downloads\Ficha_Tecnica_CMV_Artesanos_Burger.xlsx"
ABA = "Sub-Receitas"
LOJA = "Hamburgueria Artesanos"

# Título do bloco na planilha (normalizado) -> insumo-mistura do cadastro
MISTURAS = {
    "tempero smash": "Tempero Smash (mistura)",
    "tempero batata": "Tempero Batata (mistura)",
    "maionese da casa": "Maionese da Casa (caseira)",
    "molho especial": "Molho especial",
    "massa dadinho de tapioca": "Massa Dadinho de Tapioca (pronta)",
    "cebola crispy": "Cebola crispy",
    "bacon empanado": "Bacon empanado",
    "sweet barbecue": "Sweet barbecue",
    "cebola caramelizada": "Cebola caramelizada",
}
# Vão ao fogo: o rendimento real é menor que a soma dos ingredientes.
RENDIMENTO_A_CONFERIR = {"cebola crispy", "bacon empanado", "cebola caramelizada"}

# Ingrediente que precisa ser criado com um nome mais limpo que o da
# planilha. O nome da planilha continua casando com ele (resolver tira o
# parêntese do fim).
CRIAR_COMO = {
    "cebola in natura para empanar": "Cebola",
    "molho barbecue base": "Molho barbecue",
}

LINHAS_DE_TOTAL = {"total da receita", "custo por kg"}

# A receita mede de um jeito que o estoque não conta: "Limão (suco)" pede
# gramas de suco, e o estoque conta a fruta (por unidade ou por quilo).
# Converter grama de suco em grama de limão erraria o consumo, então a
# quantidade entra em branco até alguém dizer quanto suco rende.
QUANTIDADE_A_CONFERIR = {"limao suco"}

_PALAVRAS_VAZIAS = {"de", "da", "do", "com", "sem", "para", "em", "e"}


def _ler_sub_receitas(caminho):
    """[{"titulo", "rendimento", "ingredientes": [{"nome", "unidade", "quantidade", "custo"}]}]"""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if ABA not in wb.sheetnames:
        raise SystemExit(f"Aba {ABA!r} não encontrada. Abas: {wb.sheetnames}")
    receitas, atual = [], None
    for linha in wb[ABA].iter_rows(min_row=4, values_only=True):
        nome, unidade, quantidade, custo = (list(linha) + [None] * 4)[:4]
        nome = str(nome).strip() if nome else ""
        normal = _normalizar_nome_insumo(nome)
        if not nome or normal == "insumo base":
            continue
        if normal == "total da receita":
            atual["rendimento"] = quantidade
            continue
        if normal in LINHAS_DE_TOTAL:
            continue
        if not isinstance(quantidade, (int, float)):  # título de um bloco novo
            atual = {"titulo": normal, "rendimento": None, "ingredientes": []}
            receitas.append(atual)
            continue
        atual["ingredientes"].append({
            "nome": nome, "unidade": _unidade(unidade), "quantidade": float(quantidade),
            "custo": float(custo) if isinstance(custo, (int, float)) else None,
        })
    return receitas


def _localizar_mistura(titulo, cadastro):
    """Insumo-mistura de um bloco da planilha (título normalizado): o nome
    que o sistema deu a ela ("Tempero Batata (mistura)") ou o próprio
    título — em produção ela pode se chamar só "Tempero Batata". Serve
    também pra mistura usada dentro de outra, então as duas acham o mesmo
    insumo."""
    nome = MISTURAS.get(titulo)
    return localizar_insumo([nome, titulo], cadastro) if nome else None


def _parecidos(nome, cadastro, limite=3):
    """Só pro relatório, nunca pra decidir: insumos do cadastro com alguma
    palavra em comum ("Açúcar" -> "Açúcar Cristal 5Kg"). Deixa à vista o
    insumo que ia ser criado repetido."""
    palavras = set(nome_base(nome).split()) - _PALAVRAS_VAZIAS
    em_comum = {i["nome"]: len(palavras & set(nome_base(i["nome"]).split())) for i in cadastro.values()}
    return sorted((n for n, qtd in em_comum.items() if qtd), key=lambda n: (-em_comum[n], n))[:limite]


def _motivo_em_branco(ingrediente, insumo):
    if _normalizar_nome_insumo(ingrediente["nome"]) in QUANTIDADE_A_CONFERIR:
        return "a receita pede suco e o estoque conta a fruta: preencha quanto vai na tela"
    if _unidade(insumo["unidade_medida"]) not in ("g", "kg", "ml", "l"):
        return f"cadastre quanto tem em 1 {insumo['unidade_medida']} de {insumo['nome']!r}"
    return "as unidades não convertem"


def importar(caminho=PLANILHA_PADRAO, aplicar=False):
    inicializar_banco()
    receitas = _ler_sub_receitas(caminho)
    ja_tem_receita = set(mapa_receita_insumo())
    with conexao() as conn:
        cadastro = {
            _normalizar_nome_insumo(l["nome"]): dict(l)
            for l in conn.execute(
                "SELECT id, nome, unidade_medida, conteudo_por_unidade, unidade_conteudo FROM insumo"
            ).fetchall()
        }

    # criar: nome normalizado -> insumo que ainda não existe. A receita
    # aponta pro nome normalizado, que na hora de gravar é o do cadastro ou o
    # do insumo recém-criado.
    plano, criar, avisos = [], {}, []
    print(f"=== {len(receitas)} sub-receitas na planilha ===")
    for receita in receitas:
        mistura = _localizar_mistura(receita["titulo"], cadastro)
        if not mistura:
            print(f"   ! {receita['titulo']!r}: sem insumo correspondente no cadastro — pulada")
            continue
        if mistura["id"] in ja_tem_receita:
            print(f"   {mistura['nome']:34} já tem receita — mantida")
            continue

        # A planilha soma g e ml como se fossem a mesma coisa (densidade ≈ 1,
        # mesma suposição do import de custo); o rendimento vai pra unidade
        # da mistura (Massa de Dadinho é contada em kg).
        rendimento = quantidade_para_insumo(receita["rendimento"], "g", mistura)
        if not rendimento:
            print(f"   ! {mistura['nome']}: rendimento {receita['rendimento']!r} não converte pra "
                  f"{mistura['unidade_medida']!r} — pulada")
            continue

        # Ingrediente que é outra mistura ("Maionese da Casa" dentro do Molho
        # Especial) é a mistura cadastrada — nunca um insumo novo "comprado
        # pronto", senão a cascata quebraria. Sem ela no cadastro, a receita
        # espera (e nada é criado por causa dela).
        falta_mistura = next((
            i["nome"] for i in receita["ingredientes"]
            if _normalizar_nome_insumo(i["nome"]) in MISTURAS
            and _localizar_mistura(_normalizar_nome_insumo(i["nome"]), cadastro) is None
        ), None)
        if falta_mistura:
            print(f"   ! {mistura['nome']}: leva {falta_mistura!r}, que é mistura e não está no cadastro — pulada")
            continue

        linhas, textos, fora = {}, [], []
        for ingrediente in receita["ingredientes"]:
            normal = _normalizar_nome_insumo(ingrediente["nome"])
            if normal in MISTURAS:
                insumo = _localizar_mistura(normal, cadastro)
            else:
                nome_novo = CRIAR_COMO.get(normal, ingrediente["nome"])
                insumo = localizar_insumo([ingrediente["nome"], nome_novo], cadastro)
                if insumo is None:
                    insumo = criar.setdefault(_normalizar_nome_insumo(nome_novo), {
                        "nome": nome_novo, "unidade_medida": ingrediente["unidade"],
                        "custo": ingrediente["custo"], "novo": True,
                    })
            # Quantidade que não converte entra em branco, não fica de fora:
            # sem a linha, a receita pareceria completa e o custo sairia menor
            # que o real. Em branco, a receita fica incompleta (vale o custo
            # digitado) e a tela mostra a linha esperando a quantidade.
            quantidade = (None if normal in QUANTIDADE_A_CONFERIR
                          else quantidade_para_insumo(ingrediente["quantidade"], ingrediente["unidade"], insumo))
            chave = _normalizar_nome_insumo(insumo["nome"])
            if chave in linhas:
                # Dois nomes da planilha que são o mesmo insumo: soma.
                anterior = linhas[chave]
                quantidade = None if anterior is None or quantidade is None else anterior + quantidade
            linhas[chave] = quantidade

            origem = "" if chave == normal else f"{ingrediente['nome']} → "
            novo = " (novo)" if insumo.get("novo") else ""
            if quantidade is None:
                textos.append(f"{origem}{insumo['nome']}{novo} (em branco)")
                fora.append(f"{ingrediente['nome']} ({ingrediente['unidade']} na planilha, "
                            f"{insumo['unidade_medida']} no cadastro): {_motivo_em_branco(ingrediente, insumo)}")
            else:
                textos.append(f"{origem}{insumo['nome']}{novo} {quantidade:g} {insumo['unidade_medida']}")

        plano.append((mistura, rendimento, list(linhas.items())))
        conferir = "  ⚠ rende a soma dos ingredientes — conferir" if receita["titulo"] in RENDIMENTO_A_CONFERIR else ""
        print(f"   {mistura['nome']:34} rende {rendimento:g} {mistura['unidade_medida']}{conferir}")
        print(f"      {', '.join(textos)}")
        for motivo in fora:
            print(f"      ! quantidade em branco — {motivo}")
            avisos.append(f"{mistura['nome']}: {motivo}")

    if criar:
        print(f"\n=== Ingredientes novos ({len(criar)}) — vinculados ao {LOJA} ===")
        for info in criar.values():
            custo = f"R$ {info['custo']}/{info['unidade_medida']}" if info["custo"] is not None else "sem custo"
            parecidos = _parecidos(info["nome"], cadastro)
            dica = f"  — no cadastro já tem: {', '.join(parecidos)}" if parecidos else ""
            print(f"   {info['nome']:28} {info['unidade_medida']:<3} {custo}{dica}")
        print(f"\n   ⚠ {len(criar)} insumo(s) vão ser CRIADOS: {', '.join(i['nome'] for i in criar.values())}.")
        print("     Se algum já existe com outro nome, NÃO aplique — mande este relatório pra conferir.")

    if not aplicar:
        print("\nSimulação. Rode de novo com --apply pra gravar.")
        return

    agora = datetime.now().isoformat()
    for chave, info in criar.items():
        insumo_id = criar_insumo(info["nome"], "Ingrediente", info["unidade_medida"], [LOJA])
        if info["custo"] is not None:
            with conexao() as conn:
                conn.execute("UPDATE insumo SET custo_referencia = ? WHERE id = ?", (info["custo"], insumo_id))
        cadastro[chave] = {"id": insumo_id, "nome": info["nome"], "unidade_medida": info["unidade_medida"]}

    for mistura, rendimento, linhas in plano:
        ingredientes = [{"insumoId": cadastro[chave]["id"], "quantidade": q} for chave, q in linhas]
        definir_receita_insumo(mistura["id"], rendimento, ingredientes)
        # Ingrediente que já existia de outra loja (o Sal da Tradiça) passa a
        # ser usado no Artesanos também: aparece no estoque de lá.
        with conexao() as conn:
            for ingrediente in ingredientes:
                conn.execute("INSERT OR IGNORE INTO insumo_loja (insumo_id, loja) VALUES (?, ?)",
                             (ingrediente["insumoId"], LOJA))
                conn.execute(
                    "INSERT OR IGNORE INTO estoque_insumo (insumo_id, loja, quantidade_atual, estoque_minimo, atualizado_em) "
                    "VALUES (?, ?, 0, 0, ?)", (ingrediente["insumoId"], LOJA, agora))

    print(f"\nReceitas gravadas: {len(plano)}. Ingredientes criados: {len(criar)}.")
    for aviso in avisos:
        print(f"Pendente: {aviso}")


if __name__ == "__main__":
    argumentos = [a for a in sys.argv[1:] if a != "--apply"]
    importar(argumentos[0] if argumentos else PLANILHA_PADRAO, aplicar="--apply" in sys.argv)
