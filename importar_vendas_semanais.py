# -*- coding: utf-8 -*-
"""Importa o histórico semanal por canal da planilha do chefe da Julia
("Cópia de RELATÓRIO VENDAS SEMANAL.xlsx") pra tabela
faturamento_canal_semanal, que alimenta a tela "Vendas Semanais".

Uso:
  python importar_vendas_semanais.py            # simulação, não grava nada
  python importar_vendas_semanais.py --apply    # grava de verdade

Dois detalhes da planilha que o script resolve:

1. **Os períodos não têm ano** ("31/03 a 06/04"), e cada aba emenda mais de
   um ciclo anual em sequência. O ano é inferido andando de trás pra frente:
   a última linha da aba é a semana mais recente que já terminou, e cada vez
   que o mês SOBE ao voltar (jan → dez) o ano cai em 1.

2. **A coluna TOTAL da aba do Açaí não soma o 99 Food** (a coluna deve ter
   sido criada depois e a fórmula não foi refeita). O script grava canal por
   canal e deixa o total ser somado na hora de exibir, então o total do
   sistema vai ficar MAIOR que o da planilha nessas semanas — está certo,
   mas é uma diferença visível e proposital.
"""
import re
import sys
from datetime import date, datetime

import openpyxl

from backend.armazenamento import (
    inicializar_banco,
    salvar_faturamento_canal_semanal_se_ausente,
)

# O console do Windows abre em cp1252 e engasga com "Tradiça"/"Açaí".
try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass

PLANILHA = r"C:\Users\Guilherme\Downloads\Cópia de RELATÓRIO VENDAS SEMANAL.xlsx"

# Aba da planilha -> loja como o AdmFood chama
LOJA_POR_ABA = {
    "ARTESANOS": "Hamburgueria Artesanos",
    "TRADIÇA ZN": "Tradiça ZN",
    "SIMUS": "Tradiça Simus",
    "AÇAÍ NALATA": "Açaí Na Lata",
}

# Coluna da planilha (índice 0) -> canal, no mesmo vocabulário de
# NOMES_CANAL_REDE em app.py ("portal" é a venda presencial/balcão).
CANAL_POR_COLUNA = {1: "ifood", 2: "catalog", 3: "food99", 4: "portal"}

PRIMEIRA_LINHA_DADOS = 4

# "31/03 a 06/04", "31/03 á 05/04", "17/03 – 22/03", "25/05 31/05" — o
# separador varia (inclusive caractere não-ASCII), então aceita qualquer
# coisa que não seja dígito entre as duas datas.
PERIODO = re.compile(r"^\s*(\d{1,2})/(\d{1,2})\s*\D*\s*(\d{1,2})/(\d{1,2})\s*$")


def _numero(valor):
    if valor is None or valor == "":
        return None
    try:
        return round(float(valor), 2)
    except (TypeError, ValueError):
        return None


def _ler_aba(ws):
    """Linhas cruas válidas da aba: (dia_ini, mes_ini, dia_fim, mes_fim, canais)."""
    linhas = []
    ignoradas = []
    for numero_linha, r in enumerate(ws.iter_rows(min_row=PRIMEIRA_LINHA_DADOS, values_only=True), PRIMEIRA_LINHA_DADOS):
        periodo = str(r[0]).strip() if r[0] is not None else ""
        canais = {canal: _numero(r[coluna]) for coluna, canal in CANAL_POR_COLUNA.items()}
        if not periodo and all(v is None for v in canais.values()):
            continue
        casado = PERIODO.match(periodo)
        if not casado:
            ignoradas.append((numero_linha, f"{periodo!r} — período ilegível"))
            continue
        if all(v is None for v in canais.values()):
            ignoradas.append((numero_linha, f"{periodo!r} — sem nenhum valor"))
            continue
        dia_i, mes_i, dia_f, mes_f = (int(g) for g in casado.groups())
        # Dia/mês fora da faixa quer dizer célula digitada errada (ex:
        # "06/0 a 12/04", faltando um dígito no mês). Não dá pra adivinhar
        # sem inventar dado, então a linha fica de fora e é reportada — é
        # só corrigir a planilha e rodar de novo, a importação não duplica.
        if not (1 <= mes_i <= 12 and 1 <= mes_f <= 12 and 1 <= dia_i <= 31 and 1 <= dia_f <= 31):
            ignoradas.append((numero_linha, f"{periodo!r} — data inválida, corrija a planilha"))
            continue
        linhas.append((dia_i, mes_i, dia_f, mes_f, canais))
    return linhas, ignoradas


def _atribuir_anos(linhas, hoje):
    """A planilha não traz ano. Como as linhas estão em ordem cronológica,
    ancora a ÚLTIMA no ano mais recente em que ela já terminou e caminha de
    trás pra frente, tirando um ano toda vez que o mês sobe ao voltar
    (fevereiro visto logo antes de janeiro = virada de ano)."""
    if not linhas:
        return []

    _, _, _, mes_f_ultima = linhas[-1][:4]
    ano = hoje.year
    if mes_f_ultima > hoje.month:
        ano -= 1  # a última semana da planilha ainda não chegou neste ano

    datados = []
    mes_anterior = None
    for dia_i, mes_i, dia_f, mes_f, canais in reversed(linhas):
        if mes_anterior is not None and mes_i > mes_anterior:
            ano -= 1
        inicio = date(ano, mes_i, dia_i)
        # Semana que atravessa o ano novo ("29/12 a 04/01") termina no ano seguinte.
        fim = date(ano + 1 if mes_f < mes_i else ano, mes_f, dia_f)
        datados.append((inicio, fim, canais))
        mes_anterior = mes_i
    datados.reverse()
    return datados


def importar(aplicar=False):
    inicializar_banco()
    wb = openpyxl.load_workbook(PLANILHA, data_only=True)
    hoje = datetime.now().date()

    total_gravado = total_existente = 0
    for aba, loja in LOJA_POR_ABA.items():
        if aba not in wb.sheetnames:
            print(f"\n### {aba}: aba não encontrada na planilha — pulando")
            continue

        linhas, ignoradas = _ler_aba(wb[aba])
        datadas = _atribuir_anos(linhas, hoje)

        # Semana que dura muito mais que uma semana é dígito trocado na
        # planilha (ex: "02/06 á 07/09" no meio de junho, "04/10 á 10/11"
        # entre 28/10 e 11/11). Gravar isso arquivaria o faturamento na
        # semana errada, então fica de fora e é reportado.
        semanas, suspeitas = [], []
        for inicio, fim, canais in datadas:
            if not 1 <= (fim - inicio).days <= 8:
                suspeitas.append((inicio, fim))
            else:
                semanas.append((inicio, fim, canais))

        print(f"\n### {aba} → {loja}")
        for inicio, fim in suspeitas:
            print(f"  ignorada: {inicio} a {fim} — período de {(fim - inicio).days} dias, dígito trocado na planilha")
        if not semanas:
            print("  nenhuma semana válida")
            continue
        print(f"  {len(semanas)} semanas: {semanas[0][0]} até {semanas[-1][1]}")
        for numero_linha, texto in ignoradas:
            print(f"  ignorada (linha {numero_linha}): {texto!r}")

        for inicio, fim, canais in semanas:
            valores = {c: v for c, v in canais.items() if v is not None}
            if aplicar:
                for canal, valor in valores.items():
                    if salvar_faturamento_canal_semanal_se_ausente(
                        loja, inicio.isoformat(), fim.isoformat(), canal, valor
                    ):
                        total_gravado += 1
                    else:
                        total_existente += 1
            else:
                total_gravado += len(valores)

        amostra = semanas[:2] + semanas[-2:] if len(semanas) > 4 else semanas
        for inicio, fim, canais in amostra:
            resumo = "  ".join(
                f"{c}={v:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
                for c, v in canais.items() if v is not None
            )
            print(f"    {inicio} a {fim}: {resumo}")

    print()
    if aplicar:
        print(f"Gravadas {total_gravado} linhas. {total_existente} já existiam e foram mantidas como estavam.")
    else:
        print(f"Simulação: gravaria {total_gravado} linhas (canal × semana).")
        print("Confira os números acima e rode com --apply pra gravar:")
        print("  python importar_vendas_semanais.py --apply")


if __name__ == "__main__":
    importar(aplicar="--apply" in sys.argv)
