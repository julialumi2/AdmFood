# -*- coding: utf-8 -*-
"""Importa o custo por unidade de cada insumo da aba "Insumos" da planilha
de CMV do chefe (Ficha_Tecnica_CMV_Artesanos_Burger.xlsx).

Pra que serve: sem custo de insumo, o CMV de produto nenhum pode ser
calculado, e a Curva ABC de Cardápio não consegue classificar ninguém — o
sistema só saberia volume e receita. O custo da compra recebida é melhor
(é o dinheiro que saiu de verdade), mas até existir uma compra registrada
esse é o número que a casa usa.

Grava em `insumo.custo_referencia`, que é o último da fila em
`_mapa_preco_insumo`: assim que uma cotação ou uma compra recebida entrar
pro mesmo insumo, ela passa na frente sozinha.

Uso:
  python importar_custos_insumo.py            # simulação
  python importar_custos_insumo.py --apply    # grava
"""
import sys

import openpyxl

from backend.armazenamento import inicializar_banco, conexao, _normalizar_nome_insumo
from backend.nomes_insumo import resolver

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, OSError):
    pass

PLANILHA_PADRAO = r"C:\Users\Guilherme\Downloads\Ficha_Tecnica_CMV_Artesanos_Burger.xlsx"
ABA = "Insumos"
PRIMEIRA_LINHA = 5
COL_NOME, COL_UNIDADE, COL_CUSTO = 0, 1, 2

# Como cada unidade é escrita nas duas fontes.
_APELIDOS = {
    "un": "un", "unid": "un", "und": "un", "unidade": "un",
    "kg": "kg", "quilo": "kg", "g": "g", "grama": "g", "gr": "g",
    "l": "l", "lt": "l", "litro": "l", "ml": "ml",
    "fd": "fd", "cx": "cx", "pct": "pct", "pc": "pc",
}

# Conversões que dão pra fazer sem assumir nada: preço por quilo vira preço
# por grama dividindo por mil. Já kg -> un exigiria saber quanto pesa uma
# unidade, e isso o sistema não sabe — esse caso é recusado.
_FATOR = {("kg", "g"): 1000, ("g", "kg"): 0.001, ("l", "ml"): 1000, ("ml", "l"): 0.001}

# Peso <-> volume é a única conversão aqui que assume alguma coisa: que 1 g ≈
# 1 ml, ou seja, densidade perto da água. Vale pros molhos e sucos da casa (é
# o caso de todos os insumos que caem aqui hoje). Óleo daria ~8% de
# diferença — pouco perto da alternativa, que é ficar sem custo nenhum e o
# produto inteiro sumir da Curva ABC. A suposição é impressa toda vez que é
# usada. Os pares com kg e litro são a mesma suposição em outra escala: o
# catálogo da VMarket conta líquido em kg ("Molho Inglês", "Oleo De Soja"),
# e a receita pede ml.
_FATOR_COM_SUPOSICAO = {
    ("g", "ml"): 1, ("ml", "g"): 1,
    ("kg", "l"): 1, ("l", "kg"): 1,
    ("kg", "ml"): 1000, ("ml", "kg"): 0.001,
    ("l", "g"): 1000, ("g", "l"): 0.001,
}


def _unidade(bruta):
    return _APELIDOS.get(str(bruta or "").strip().lower(), str(bruta or "").strip().lower())


def _converter_custo(custo, unidade_planilha, unidade_cadastro):
    """Devolve (custo_na_unidade_do_cadastro, observação) ou (None, motivo)
    quando as unidades não são conversíveis. Gravar um preço por quilo num
    insumo contado por unidade faria cada hambúrguer custar o preço de um
    quilo de carne — foi exatamente o que aconteceu antes desta checagem."""
    origem, destino = _unidade(unidade_planilha), _unidade(unidade_cadastro)
    if not origem or not destino or origem == destino:
        return custo, None
    fator = _FATOR.get((origem, destino))
    if fator:
        return round(custo / fator, 6), f"convertido de /{origem} pra /{destino}"
    fator = _FATOR_COM_SUPOSICAO.get((origem, destino))
    if fator:
        return round(custo / fator, 6), f"de /{origem} pra /{destino} assumindo 1 g ≈ 1 ml"
    return None, f"unidade incompatível: planilha em {origem!r}, cadastro em {destino!r}"


def importar(caminho, aplicar=False):
    inicializar_banco()
    wb = openpyxl.load_workbook(caminho, data_only=True)
    if ABA not in wb.sheetnames:
        raise SystemExit(f"Aba {ABA!r} não encontrada. Abas: {wb.sheetnames}")

    da_planilha = {}
    for linha in wb[ABA].iter_rows(min_row=PRIMEIRA_LINHA, values_only=True):
        nome = str(linha[COL_NOME]).strip() if linha[COL_NOME] else ""
        custo = linha[COL_CUSTO]
        # Linha de seção ("PÃES") vem sem custo — não é insumo.
        if not nome or not isinstance(custo, (int, float)):
            continue
        da_planilha[_normalizar_nome_insumo(nome)] = {
            "nome": nome,
            "custo": round(float(custo), 4),
            "unidade": linha[COL_UNIDADE],
        }

    with conexao() as conn:
        cadastrados = {
            _normalizar_nome_insumo(l["nome"]): dict(l)
            for l in conn.execute("SELECT id, nome, unidade_medida, custo_referencia FROM insumo").fetchall()
        }

    casados, novos_valores, sem_cadastro, incompativeis = [], 0, [], []
    for info in da_planilha.values():
        alvo = resolver(info["nome"], cadastrados)
        if not alvo:
            sem_cadastro.append(info["nome"])
            continue
        custo, nota = _converter_custo(info["custo"], info["unidade"], alvo["unidade_medida"])
        if custo is None:
            incompativeis.append((alvo, info, nota))
            continue
        info = {**info, "custo": custo, "nota": nota}
        casados.append((alvo, info))
        if alvo["custo_referencia"] != custo:
            novos_valores += 1

    print(f"Planilha: {len(da_planilha)} insumos com custo")
    print(f"Casaram: {len(casados)}  |  sem cadastro: {len(sem_cadastro)}  |  unidade incompatível: {len(incompativeis)}")

    if incompativeis:
        print("\nRECUSADOS — a unidade não bate, e converter exigiria assumir peso/volume por unidade.")
        print("Ajuste a unidade do insumo no sistema (ou na planilha) e rode de novo:")
        for alvo, info, nota in incompativeis:
            atual = "—" if alvo["custo_referencia"] is None else f"R$ {alvo['custo_referencia']}"
            print(f"   {info['nome'][:30]:30} -> {alvo['nome'][:30]:30} {nota}  (custo hoje: {atual})")
    if sem_cadastro:
        print("\nSem cadastro no sistema (não vão entrar — cadastre o insumo antes, se for usar):")
        for nome in sem_cadastro:
            print(f"   {nome}")

    print(f"\nValores a gravar/atualizar: {novos_valores}")
    for alvo, info in casados[:8]:
        antes = "—" if alvo["custo_referencia"] is None else f"R$ {alvo['custo_referencia']}"
        print(f"   {alvo['nome'][:34]:34} {antes:>10}  ->  R$ {info['custo']}")
    if len(casados) > 8:
        print(f"   ... e mais {len(casados) - 8}")

    if not aplicar:
        print("\nSimulação. Rode de novo com --apply pra gravar.")
        return

    limpos = [a for a, _i, _n in incompativeis if a["custo_referencia"] is not None]
    with conexao() as conn:
        for alvo, info in casados:
            conn.execute("UPDATE insumo SET custo_referencia = ? WHERE id = ?", (info["custo"], alvo["id"]))
        # Custo de unidade incompatível é apagado, não deixado como está:
        # esta coluna só é preenchida por este script, então valor que
        # sobrou ali veio de uma rodada anterior sem a checagem — é
        # justamente o número errado que se quer tirar de circulação.
        for alvo in limpos:
            conn.execute("UPDATE insumo SET custo_referencia = NULL WHERE id = ?", (alvo["id"],))

    print(f"\nGravados {len(casados)} custos.")
    if limpos:
        print(f"Apagados {len(limpos)} custos que tinham sido gravados com unidade incompatível:")
        for alvo in limpos:
            print(f"   {alvo['nome']}")


if __name__ == "__main__":
    argumentos = [a for a in sys.argv[1:] if a != "--apply"]
    importar(argumentos[0] if argumentos else PLANILHA_PADRAO, aplicar="--apply" in sys.argv)
